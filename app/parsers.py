"""Извлечение содержимого из документов разных форматов.

Каждый файл приводится к одному из двух представлений:
  * TextDoc  — список строк (Word, PDF, PowerPoint, текстовые файлы);
  * TableDoc — набор листов с ячейками (Excel, CSV).
Дальше их сравнивает app.differ.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field

from . import config


class UnsupportedFormat(ValueError):
    pass


@dataclass
class TextDoc:
    lines: list[str]
    kind: str = "text"


@dataclass
class TableDoc:
    sheets: dict[str, list[list[str]]] = field(default_factory=dict)
    kind: str = "table"

    def to_text(self) -> TextDoc:
        lines: list[str] = []
        for name, rows in self.sheets.items():
            lines.append(f"## Лист: {name}")
            lines.extend(" | ".join(r) for r in rows)
        return TextDoc(lines)


TEXT_EXT = {".txt", ".md", ".json", ".xml", ".yaml", ".yml", ".html", ".htm",
            ".ini", ".cfg", ".conf", ".log", ".sql", ".py", ".js", ".ts", ".java", ".go"}
# Форматы, которые можно открыть только через LibreOffice (если он есть на сервере)
SOFFICE_EXT = {".doc": "docx", ".rtf": "docx", ".odt": "docx",
               ".ods": "xlsx", ".ppt": "pptx", ".odp": "pptx"}
NATIVE_EXT = {".docx", ".docm", ".xlsx", ".xlsm", ".xls", ".csv", ".tsv", ".pdf", ".pptx"} | TEXT_EXT


def supported_extensions() -> list[str]:
    exts = set(NATIVE_EXT)
    if soffice_path():
        exts |= set(SOFFICE_EXT)
    return sorted(exts)


def soffice_path() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


# ---------------------------------------------------------------- helpers

def cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "ИСТИНА" if v else "ЛОЖЬ"
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return repr(v)
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    return str(v).strip()


def _trim(rows: list[list[str]]) -> list[list[str]]:
    """Убирает пустые хвосты строк и пустые строки в конце листа."""
    rows = [list(r) for r in rows]
    for r in rows:
        while r and r[-1] == "":
            r.pop()
    while rows and not rows[-1]:
        rows.pop()
    return rows


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


# ---------------------------------------------------------------- parsers

def parse_docx(data: bytes) -> TextDoc:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    lines: list[str] = []

    def add_table(tbl: Table) -> None:
        for row in tbl.rows:
            cells, prev = [], None
            for c in row.cells:
                # объединённые ячейки python-docx возвращает повторно
                if prev is not None and c._tc is prev:
                    continue
                prev = c._tc
                cells.append(" ".join(p.text.strip() for p in c.paragraphs if p.text.strip()))
            if any(cells):
                lines.append("| " + " | ".join(cells) + " |")

    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.lower().startswith(("heading", "заголовок")):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                text = "#" * min(int(level), 6) + " " + text
            lines.append(text)
        elif tag == "tbl":
            add_table(Table(el, doc))
    return TextDoc(lines)


def parse_pdf(data: bytes) -> TextDoc:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    lines: list[str] = []
    for i, page in enumerate(reader.pages, 1):
        lines.append(f"— Страница {i} —")
        text = page.extract_text() or ""
        lines.extend(ln.strip() for ln in text.splitlines() if ln.strip())
    return TextDoc(lines)


def parse_pptx(data: bytes) -> TextDoc:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    lines: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"— Слайд {i} —")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        lines.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[Заметки] {notes}")
    return TextDoc(lines)


def parse_xlsx(data: bytes) -> TableDoc:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    doc = TableDoc()
    for ws in wb.worksheets:
        rows = [[cell_to_str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        doc.sheets[ws.title] = _trim(rows)
    wb.close()
    return doc


def parse_xls(data: bytes) -> TableDoc:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    doc = TableDoc()
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    v = bool(v)
                row.append(cell_to_str(v))
            rows.append(row)
        doc.sheets[sh.name] = _trim(rows)
    return doc


def parse_csv(data: bytes, ext: str) -> TableDoc:
    text = decode_text(data)
    if ext == ".tsv":
        delim = "\t"
    else:
        try:
            delim = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delim = ";" if text[:4096].count(";") > text[:4096].count(",") else ","
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
    return TableDoc({"CSV": _trim(rows)})


def parse_text(data: bytes) -> TextDoc:
    return TextDoc([ln.rstrip() for ln in decode_text(data).splitlines()])


def convert_with_soffice(data: bytes, ext: str) -> tuple[bytes, str]:
    exe = soffice_path()
    target = SOFFICE_EXT[ext]
    if not exe:
        raise UnsupportedFormat(
            f"Формат {ext} поддерживается только при установленном LibreOffice. "
            f"Сохраните файл как .{target} или соберите образ с WITH_LIBREOFFICE=1."
        )
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "input" + ext)
        with open(src, "wb") as f:
            f.write(data)
        env = dict(os.environ, HOME=tmp)
        subprocess.run(
            [exe, "--headless", "--norestore", "--convert-to", target, "--outdir", tmp, src],
            check=True, capture_output=True, timeout=config.SOFFICE_TIMEOUT, env=env,
        )
        with open(os.path.join(tmp, "input." + target), "rb") as f:
            return f.read(), "." + target


def parse(data: bytes, filename: str) -> TextDoc | TableDoc:
    ext = os.path.splitext(filename.lower())[1]
    if ext in SOFFICE_EXT:
        data, ext = convert_with_soffice(data, ext)
    try:
        if ext in (".docx", ".docm"):
            return parse_docx(data)
        if ext in (".xlsx", ".xlsm"):
            return parse_xlsx(data)
        if ext == ".xls":
            return parse_xls(data)
        if ext in (".csv", ".tsv"):
            return parse_csv(data, ext)
        if ext == ".pdf":
            return parse_pdf(data)
        if ext == ".pptx":
            return parse_pptx(data)
        if ext in TEXT_EXT or ext == "":
            return parse_text(data)
    except UnsupportedFormat:
        raise
    except Exception as e:  # битый файл, неверное расширение и т.п.
        raise UnsupportedFormat(f"Не удалось прочитать «{filename}»: {e}") from e
    raise UnsupportedFormat(f"Формат «{ext}» не поддерживается")
