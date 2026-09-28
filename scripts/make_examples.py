"""Создаёт пары демо-файлов в examples/ для быстрой проверки."""
from pathlib import Path

from docx import Document
from openpyxl import Workbook

OUT = Path(__file__).resolve().parent.parent / "examples"
OUT.mkdir(exist_ok=True)


def contract(path, price, days, extra):
    d = Document()
    d.add_heading("Договор поставки № 125/24", 1)
    d.add_paragraph("г. Москва")
    d.add_heading("1. Предмет договора", 2)
    d.add_paragraph("Поставщик обязуется поставить товар, а Покупатель — принять и оплатить его.")
    d.add_heading("2. Цена и порядок расчётов", 2)
    d.add_paragraph(f"Общая стоимость товара составляет {price} рублей, включая НДС 20%.")
    d.add_paragraph(f"Оплата производится в течение {days} банковских дней с даты поставки.")
    if extra:
        d.add_paragraph("Покупатель вправе удержать неустойку из суммы оплаты.")
    d.add_heading("3. Ответственность сторон", 2)
    d.add_paragraph("За просрочку поставки Поставщик уплачивает пени в размере 0,1% за каждый день.")
    t = d.add_table(rows=3, cols=3)
    for r, row in enumerate([["Товар", "Кол-во", "Цена"], ["Молоко 3,2%", "1000", "89"],
                             ["Хлеб белый", "500", "45" if not extra else "49"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.save(path)


def pricelist(path, changed):
    wb = Workbook()
    ws = wb.active
    ws.title = "Прайс"
    ws.append(["Артикул", "Наименование", "Цена", "Остаток"])
    items = [("1001", "Молоко 3,2% 1л", 89, 120), ("1002", "Кефир 1% 0,9л", 79, 80),
             ("1003", "Хлеб белый", 45, 300), ("1004", "Сыр Российский 200г", 199, 45),
             ("1005", "Масло сливочное 82%", 229, 60)]
    if changed:
        items[1] = ("1002", "Кефир 1% 0,9л", 84, 80)
        items.insert(3, ("1006", "Йогурт питьевой", 65, 150))
        del items[-1]
    for it in items:
        ws.append(list(it))
    ws2 = wb.create_sheet("Поставщики")
    ws2.append(["Поставщик", "ИНН"])
    ws2.append(["ООО «Молочный край»", "7701234567"])
    if changed:
        wb.create_sheet("Акции").append(["Артикул", "Скидка"])
    wb.save(path)


contract(OUT / "contract_v1.docx", "1 250 000", 30, False)
contract(OUT / "contract_v2.docx", "1 310 000", 45, True)
pricelist(OUT / "price_v1.xlsx", False)
pricelist(OUT / "price_v2.xlsx", True)
print(f"Примеры созданы в {OUT}")
