import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.differ import Options, compare_files, diff_text, inline_diff
from app.main import app
from app.parsers import TableDoc, parse


def _docx(*paras):
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx(rows, title="Лист1"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_identical_text():
    r = diff_text(["a", "b"], ["a", "b"], Options())
    assert r["stats"]["similarity"] == 1.0
    assert all(row["op"] == "equal" for row in r["rows"])


def test_insert_delete_replace():
    r = diff_text(["один", "два", "три"], ["один", "два!", "три", "четыре"], Options())
    ops = [row["op"] for row in r["rows"]]
    assert ops == ["equal", "replace", "equal", "insert"]


def test_inline_diff_words():
    a, b = inline_diff("цена 100 рублей", "цена 120 рублей", Options())
    assert ["-", "100"] in a and ["+", "120"] in b


def test_ignore_options():
    r = diff_text(["Привет  мир"], ["привет мир"], Options(ignore_case=True, ignore_whitespace=True))
    assert r["rows"][0]["op"] == "equal"


def test_docx_compare():
    r = compare_files(_docx("Срок оплаты 30 дней", "Конец"), "a.docx",
                      _docx("Срок оплаты 45 дней", "Конец"), "b.docx")
    assert r["mode"] == "text"
    assert r["stats"]["replace"] == 1


def test_xlsx_compare_row_insert_and_cell_change():
    a = _xlsx([["id", "price"], [1, 10], [2, 20]])
    b = _xlsx([["id", "price"], [1, 10], [3, 30], [2, 25]])
    r = compare_files(a, "a.xlsx", b, "b.xlsx")
    assert r["mode"] == "table"
    sh = r["sheets"][0]
    assert sh["status"] == "changed"
    ops = [row["op"] for row in sh["rows"]]
    assert "insert" in ops
    changed = [c for row in sh["rows"] if row["op"] == "replace" for c in row["cells"] if c[0] == "~"]
    assert changed == [["~", "20", "25"]]


def test_sheets_added_removed():
    from app.differ import diff_tables
    a = TableDoc({"S1": [["x"]], "S2": [["y"]]})
    b = TableDoc({"S1": [["x"]], "S3": [["z"]]})
    r = diff_tables(a, b, Options())
    assert [s["status"] for s in r["sheets"]] == ["same", "removed", "added"]


def test_csv_semicolon():
    doc = parse("a;b\n1;2\n".encode("cp1251"), "x.csv")
    assert doc.sheets["CSV"] == [["a", "b"], ["1", "2"]]


def test_mixed_formats_fall_back_to_text():
    r = compare_files(_xlsx([["a", "b"]]), "a.xlsx", b"a | b\n", "b.txt")
    assert r["mode"] == "text"


def test_api():
    c = TestClient(app)
    assert c.get("/health").json()["status"] == "ok"
    assert "text/html" in c.get("/").headers["content-type"]
    resp = c.post("/api/compare", files={"file_a": ("a.txt", b"x\ny\n"), "file_b": ("b.txt", b"x\nz\n")})
    assert resp.status_code == 200
    assert resp.json()["stats"]["replace"] + resp.json()["stats"]["delete"] >= 1


def test_api_unsupported():
    c = TestClient(app)
    resp = c.post("/api/compare", files={"file_a": ("a.exe", b"x"), "file_b": ("b.exe", b"y")})
    assert resp.status_code == 415
